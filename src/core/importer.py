"""`Uy moliya.xlsx` dan ma'lumot ko'chirish.

Excel'dagi "Ha" belgilari shu yerda haqiqiy hisob-kitob to'lovlariga
aylanadi: agar Abbosxon 11 ta qatorda "Ha" bo'lsa, u o'sha qatorlardagi
ulushlari yig'indisini to'lagan deb yoziladi — bitta to'lov qatori bilan.
"""
from __future__ import annotations

from dataclasses import dataclass, field
from datetime import date, datetime

from core import entries

ODAMLAR = ("Fayzulloxon", "Otabek", "Abbosxon")

# Barcha varaqlar B ustunidan boshlanadi. min_col=2 bilan o'qiymiz, shuning
# uchun indeks 0 = B, 1 = C, ...
B_USTUN = 2

# "Umumiy rasxodlar" varag'idagi "<kim> to'ladimi?" ustunlari:
#   B=0 sana, C=1 summa, D=2 sabab, E=3 kim to'ladi, F=4 har biriga,
#   G=5 Abbosxon, H=6 Otabek, I=7 Fayzulloxon
TOLADI_USTUN = {"Abbosxon": 5, "Otabek": 6, "Fayzulloxon": 7}


@dataclass
class Natija:
    odam: int = 0
    kirim: int = 0
    umumiy: int = 0
    shaxsiy: int = 0
    qarz: int = 0
    hisob_kitob: int = 0
    ogohlantirish: list[str] = field(default_factory=list)

    def hisobot(self) -> str:
        q = [f"{self.odam} odam", f"{self.kirim} kirim",
             f"{self.umumiy} umumiy rasxod", f"{self.shaxsiy} shaxsiy rasxod",
             f"{self.qarz} qarz", f"{self.hisob_kitob} hisob-kitob"]
        s = "Ko'chirildi: " + ", ".join(q)
        if self.ogohlantirish:
            s += "\n\nDiqqat:\n" + "\n".join("  • " + o for o in self.ogohlantirish)
        return s


def _sana(qiymat) -> str | None:
    """Excel'dagi har xil sana formatlarini ISO ga keltiradi."""
    if qiymat is None:
        return None
    if isinstance(qiymat, datetime):
        return qiymat.date().isoformat()
    if isinstance(qiymat, date):
        return qiymat.isoformat()
    t = str(qiymat).strip()
    if not t:
        return None
    for ajratgich in (".", "/", "-"):
        if ajratgich in t:
            qismlar = t.split(ajratgich)
            if len(qismlar) == 3:
                try:
                    a, b, c = (int(x) for x in qismlar)
                except ValueError:
                    continue
                # Excel'da "8.27.2026" = oy.kun.yil
                if c > 1000:
                    oy, kun, yil = (a, b, c) if a <= 12 else (b, a, c)
                    try:
                        return date(yil, oy, kun).isoformat()
                    except ValueError:
                        return None
                if a > 1000:
                    try:
                        return date(a, b, c).isoformat()
                    except ValueError:
                        return None
    return None


def _son(qiymat) -> int:
    if qiymat is None:
        return 0
    if isinstance(qiymat, (int, float)):
        return int(round(qiymat))
    t = str(qiymat).replace(" ", "").replace(" ", "").replace(",", "")
    try:
        return int(round(float(t)))
    except ValueError:
        return 0


def _ha(qiymat) -> bool:
    return str(qiymat or "").strip().lower() in ("ha", "ha ", "yes", "1", "true")


def import_qil(db, yol: str) -> Natija:
    import openpyxl

    n = Natija()
    wb = openpyxl.load_workbook(yol, data_only=True)

    # ── odamlar ──────────────────────────────────────────────────────
    idlar: dict[str, int] = {}
    for nom in ODAMLAR:
        r = db.q1("SELECT id FROM odam WHERE nom=?", nom)
        if r:
            idlar[nom] = r["id"]
        else:
            idlar[nom] = entries.odam_qosh(db, nom)
            n.odam += 1

    # ── Kirim ────────────────────────────────────────────────────────
    if "Kirim" in wb.sheetnames:
        ws = wb["Kirim"]
        oxirgi = None
        for row in ws.iter_rows(min_col=B_USTUN, min_row=10, max_row=49, values_only=True):
            # B..E  →  indeks 0..3 (dimension B1 dan boshlanadi)
            sana, kim, summa, sabab = row[0], row[1], row[2], row[3]
            summa = _son(summa)
            kim = str(kim or "").strip()
            if not summa or kim not in idlar:
                continue
            s = _sana(sana) or oxirgi
            if not s:
                n.ogohlantirish.append(f"Kirim: sanasiz qator ({kim}, {summa}) tashlab ketildi")
                continue
            oxirgi = s
            entries.kirim_qosh(db, s, idlar[kim], summa,
                               str(sabab).strip() if sabab else None)
            n.kirim += 1

    # ── Umumiy rasxodlar ─────────────────────────────────────────────
    # Ulushlarni O'ZIMIZ bo'lamiz (butun songa), Excel'dagi F ustuniga ishonmaymiz.
    tolangan: dict[tuple[str, str], int] = {}   # (kim_toladi, kim_qaytardi) -> summa
    if "Umumiy rasxodlar" in wb.sheetnames:
        ws = wb["Umumiy rasxodlar"]
        oxirgi = None
        for row in ws.iter_rows(min_col=B_USTUN, min_row=10, max_row=39, values_only=True):
            sana, summa, sabab, kim = row[0], row[1], row[2], row[3]
            summa = _son(summa)
            kim = str(kim or "").strip()
            if not summa or kim not in idlar:
                continue
            s = _sana(sana) or oxirgi
            if not s:
                n.ogohlantirish.append(f"Umumiy rasxod {summa}: sanasiz, tashlab ketildi")
                continue
            oxirgi = s

            rid = entries.rasxod_qosh(
                db, s, str(sabab).strip() if sabab else "", summa,
                idlar[kim], umumiymi=True)
            n.umumiy += 1

            # "Ha" belgilangan odamlarning ulushini yig'amiz
            ulushlar = {r["odam_id"]: r["summa"]
                        for r in db.q("SELECT odam_id, summa FROM ulush WHERE rasxod_id=?", rid)}
            for nom, ustun in TOLADI_USTUN.items():
                if nom == kim or nom not in idlar:
                    continue          # to'lovchi o'ziga qaytarmaydi
                if ustun < len(row) and _ha(row[ustun]):
                    ulush = ulushlar.get(idlar[nom], 0)
                    if ulush:
                        kalit = (kim, nom)
                        tolangan[kalit] = tolangan.get(kalit, 0) + ulush

    # ── Shaxsiy rasxodlar (har odamning varag'i) ─────────────────────
    for nom in ODAMLAR:
        if nom not in wb.sheetnames:
            continue
        ws = wb[nom]
        for row in ws.iter_rows(min_col=B_USTUN, min_row=10, max_row=39, values_only=True):
            sana, summa, sabab = row[0], row[1], row[2]
            summa = _son(summa)
            if not summa:
                continue
            s = _sana(sana)
            if not s:
                continue
            entries.rasxod_qosh(db, s, str(sabab).strip() if sabab else "",
                                summa, idlar[nom], umumiymi=False)
            n.shaxsiy += 1

    # ── Qarz ─────────────────────────────────────────────────────────
    if "Qarz" in wb.sheetnames:
        ws = wb["Qarz"]
        oxirgi = None
        for row in ws.iter_rows(min_col=B_USTUN, min_row=10, max_row=49, values_only=True):
            sana, berdi, kimga, summa, sabab, tolandi = (
                row[0], row[1], row[2], row[3], row[4], row[5])
            summa = _son(summa)
            berdi = str(berdi or "").strip()
            kimga = str(kimga or "").strip()
            if not summa or berdi not in idlar or kimga not in idlar or berdi == kimga:
                continue
            s = _sana(sana) or oxirgi
            if not s:
                s = db.skalyar("SELECT MIN(sana) FROM kirim", birlamchi=None) or "2026-01-01"
                n.ogohlantirish.append(
                    f"Qarz {berdi}→{kimga} {summa}: Excel'da sanasi yo'q edi, {s} qo'yildi")
            oxirgi = s
            entries.qarz_qosh(db, s, idlar[berdi], idlar[kimga], summa,
                              str(sabab).strip() if sabab else None)
            n.qarz += 1
            if _ha(tolandi):
                entries.hisob_kitob_qosh(db, s, idlar[kimga], idlar[berdi], summa,
                                         "Qarz to'landi (Excel)")
                n.hisob_kitob += 1

    # ── "Ha" larni bitta hisob-kitobga aylantirish ───────────────────
    oxirgi_sana = db.skalyar("SELECT MAX(sana) FROM rasxod", birlamchi=None) \
        or date.today().isoformat()
    for (kim_toladi, kim_qaytardi), summa in sorted(tolangan.items()):
        if summa <= 0:
            continue
        entries.hisob_kitob_qosh(
            db, oxirgi_sana, idlar[kim_qaytardi], idlar[kim_toladi], summa,
            "Excel'dagi \"Ha\" belgilari bo'yicha ulush qaytarildi")
        n.hisob_kitob += 1

    return n
